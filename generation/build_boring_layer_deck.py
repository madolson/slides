"""Build "The Boring Layer" keynote on the oss_elc_EU26 conference template.

Text-only pass: one statement per slide, points in the speaker notes.
Spec: Work Vault/content/talks/The Boring Layer - What AI Is Building On, and Who Owns It.md

    python3 generation/build_boring_layer_deck.py [out.pptx]
"""
import copy
import os
import sys
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.packuri import PackURI
from pptx.parts.slide import SlideLayoutPart
from pptx.util import Emu, Pt

BASE = "/Users/matolson/Downloads/oss_elc_EU26.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "slides", "The Boring Layer.pptx")
ASSETS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets")
WHITE = "WHITE_STATEMENT"
BIG = {"AI won't say no", "Maintainers say no"}  # 44pt; every other statement is the layout's 32pt
QRS = [("qr-valkey-io.png", "valkey.io"), ("qr-trivia.png", "Trivia")]

AMPLIFYING = "Source: Amplifying.ai Valkey benchmark, 1 September 2026. 5 models, 4 repos, 100 prompts, 2,900 answers. No prompt mentions Valkey. https://chorus.aws.dev/doc/1q6EixyPkWpq/Amplifyingai---Valkey-Benchmark-Results"

# (kind, title, body lines, notes)
SLIDES = [
    ("title", "The Boring Layer: What AI Is Building On, and Who Owns It",
     ["Madelyn Olson, Valkey maintainer"], ""),
    ("statement", "Everyone is building their own software", [],
     "The cost of software that fits exactly one person or one team is collapsing. AI made the writing cheap."),
    ("statement", "1 in 6 AI agents built their own cache", [],
     "482 of 2,900 answers were custom or DIY instead of an existing store.\n" + AMPLIFYING),
    ("statement", "AI won't say no", [],
     "It builds every idea, good or bad. It never asks what problem the idea solves."),
    ("statement", "AI will make unstable infrastructure", [],
     "Ask the same question twice and you can get a different answer. In 118 of 297 model, repo and prompt combinations where Valkey was picked at least once, it was picked in only one of the two runs.\n"
     "Agents weigh the license only when the prompt raises it: Valkey is the pick in 16% of all answers and 66% of licensing prompts.\n" + AMPLIFYING),
    ("statement", "You're responsible to maintain your software", [],
     "Personalized software has no maintainer, and the AI that built it doesn't know how to be one. That leaves you."),
    ("statement", "Maintainers say no", [],
     "It's not no to contributing. It's no to bad ideas.\n"
     "Vector sets, valkey-io/valkey#2569. Redis shipped them in April 2025 and Valkey still hasn't. Madelyn, 2 September 2025: \"One thing I see missing from this thread, to my surprise, is what problems we are concretely trying to solve.\"\n"
     "Ping Xie, valkey-io/valkey#274, 2024: \"Often, removing a feature is far more challenging than adding one.\"\n"
     "557 Valkey pull requests closed without merging, 2,294 merged (GitHub search, 7 October 2026). Not every close is a deliberate no."),
    ("statement", "You should channel your inner maintainer", [],
     "Ask what problem it solves before letting the AI build it. Agents follow the constraint when you give it: when policy requires permissive licensing and neutral governance, 18 of 20 answers pick Valkey.\n" + AMPLIFYING),
    ("statement", "Review the decisions your AIs are making", [],
     "One agent picked Valkey for the wrong reason: it claimed Valkey 8 runs commands in parallel across cores. Check the reasoning, not just the pick.\n" + AMPLIFYING),
    ("statement", "AI is writing a lot of your software.\vMake sure it's building on a good foundation.", [],
     "Proven, vendor neutral open source. One owner can change the terms: Redis left BSD on 20 March 2024, and the Valkey repo was created 22 March 2024. Contribute back, and the maintainers carry your changes with you."),
    ("qr", None, [], "valkey.io: https://valkey.io. Trivia: https://d3pojgfcwxoi3q.cloudfront.net/trivia. Made with madelynolson.com/qr-generator, decoded with CoreImage to confirm."),
    ("thanks", "Thank you!", ["Madelyn Olson", "Valkey maintainer · Principal engineer at AWS"], ""),
]


def clear_slides(prs):
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        prs.part.drop_rel(sld.rId)
        ids.remove(sld)


def add_white_layout(prs):
    """Black-on-white statement layout, cloned from the template's white TITLE_AND_BODY.

    Keeps the conference branding. Title is centred and black; the body is a
    centred, bullet-free black list under it.
    """
    src = layout(prs, "TITLE_AND_BODY")
    el = copy.deepcopy(src._element)
    el.cSld.set("name", WHITE)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
          "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
    A = "{%s}" % ns["a"]

    def style(sp, off_y, cy, size):
        sp.find("p:spPr/a:xfrm/a:off", ns).set("y", str(off_y))
        sp.find("p:spPr/a:xfrm/a:ext", ns).set("cy", str(cy))
        sp.find("p:txBody/a:bodyPr", ns).set("anchor", "ctr")
        lvl1 = sp.find("p:txBody/a:lstStyle/a:lvl1pPr", ns)
        lvl1.set("algn", "ctr"); lvl1.set("marL", "0"); lvl1.set("indent", "0")
        for bu in lvl1.findall("a:buChar", ns) + lvl1.findall("a:buNone", ns):
            lvl1.remove(bu)
        old = lvl1.find("a:defRPr", ns)
        idx = list(lvl1).index(old); lvl1.remove(old)
        lvl1.insert(idx, etree.Element(A + "buNone"))
        rpr = etree.Element(A + "defRPr", sz=str(size * 100))
        fill = etree.SubElement(rpr, A + "solidFill")
        etree.SubElement(fill, A + "srgbClr", val="000000")
        lvl1.insert(idx + 1, rpr)

    for sp in el.findall(".//p:sp", ns):
        ph = sp.find("p:nvSpPr/p:nvPr/p:ph", ns)
        if ph is None:
            continue
        if ph.get("type") == "title":
            style(sp, 1450000, 1500000, 32)
        elif ph.get("idx") == "1":
            style(sp, 2950000, 1200000, 24)

    part = SlideLayoutPart(PackURI("/ppt/slideLayouts/slideLayout%d.xml" % (len(prs.slide_layouts) + 1)),
                           src.part.content_type, src.part.package, el)
    for rel in sorted(src.part.rels.values(), key=lambda r: int(r.rId[3:])):
        rid = part.relate_to(rel.target_part, rel.reltype)
        assert rid == rel.rId, (rid, rel.rId)
    master = prs.slide_master
    rid = master.part.relate_to(part, RT.SLIDE_LAYOUT)
    lst = master._element.find("p:sldLayoutIdLst", ns)
    new_id = max(int(x.get("id")) for x in lst) + 1
    etree.SubElement(lst, "{%s}sldLayoutId" % ns["p"], id=str(new_id)).set(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", rid)


def qr_slide(s, slide_w):
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx != 12:
            ph._element.getparent().remove(ph._element)
    head = s.shapes.add_textbox(Emu(311700), Emu(250000), slide_w - Emu(623400), Emu(750000))
    hp = head.text_frame.paragraphs[0]
    hp.text = "Learn more about Valkey"
    hp.alignment = PP_ALIGN.CENTER
    hp.runs[0].font.size = Pt(32)
    hp.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    size, gap, top = Emu(2300000), Emu(1000000), Emu(1100000)
    left = (slide_w - 2 * size - gap) // 2
    for i, (img, label) in enumerate(QRS):
        x = left + i * (size + gap)
        s.shapes.add_picture(f"{ASSETS}/{img}", x, top, size, size)
        tb = s.shapes.add_textbox(x, top + size + Emu(80000), size, Emu(400000))
        p = tb.text_frame.paragraphs[0]
        p.text = label
        p.alignment = PP_ALIGN.CENTER
        p.runs[0].font.size = Pt(20)
        p.runs[0].font.color.rgb = RGBColor(0, 0, 0)


def layout(prs, name):
    return next(l for l in prs.slide_layouts if l.name == name)


def main():
    prs = Presentation(BASE)
    clear_slides(prs)
    add_white_layout(prs)
    kinds = {"title": "TITLE", "thanks": "TITLE", "statement": WHITE, "list": WHITE, "logo": "Logo", "qr": WHITE}
    for kind, title, body, notes in SLIDES:
        if kind in ("statement", "list"):
            assert "," not in title, f"comma in title: {title}"
        s = prs.slides.add_slide(layout(prs, kinds[kind]))
        if kind == "qr":
            qr_slide(s, prs.slide_width)
            continue
        for ph in list(s.placeholders):
            idx = ph.placeholder_format.idx
            if idx == 0:
                ph.text_frame.text = title
                if title in BIG:
                    ph.text_frame.paragraphs[0].runs[0].font.size = Pt(44)
            elif idx == 1 and body:
                tf = ph.text_frame
                tf.text = body[0]
                for line in body[1:]:
                    tf.add_paragraph().text = line
                if kind == "thanks":
                    tf.auto_size = MSO_AUTO_SIZE.NONE
                    for i, para in enumerate(tf.paragraphs):
                        para.runs[0].font.size = Pt(20)
                        para.runs[0].font.bold = i == 0
            elif idx != 12:
                ph._element.getparent().remove(ph._element)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    qr = next(i for i, sl in enumerate(SLIDES) if sl[0] == "qr")
    prs.slides[qr].notes_slide.notes_text_frame.text = SLIDES[qr][3]
    prs.save(OUT)
    print(f"{OUT}: {len(prs.slides)} slides")


if __name__ == "__main__":
    main()
