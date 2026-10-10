#!/usr/bin/env python3
"""Dump a pptx to text: layouts, shapes, positions, runs, charts, notes.

Usage: python3 dump_deck.py "The Vertical Scaling wall.pptx" [slide_number ...]
"""
import sys
from pptx import Presentation


def walk(shapes, depth=0):
    pad = "  " * depth
    for sh in shapes:
        line = f"{pad}- [{sh.shape_type}] {sh.name!r}"
        try:
            line += f" pos=({sh.left},{sh.top}) size=({sh.width},{sh.height})"
        except Exception:
            pass
        if sh.is_placeholder:
            line += f" PH(idx={sh.placeholder_format.idx} {sh.placeholder_format.type})"
        print(line)

        if sh.shape_type == 6:  # GROUP
            walk(sh.shapes, depth + 1)
            continue

        if sh.has_text_frame:
            for i, p in enumerate(sh.text_frame.paragraphs):
                if not p.runs:
                    continue
                runs = " | ".join(
                    f"{r.text!r}(sz={r.font.size and r.font.size.pt},b={r.font.bold})"
                    for r in p.runs
                )
                print(f"{pad}    p{i}: {runs}")

        if sh.has_table:
            for row in sh.table.rows:
                print(f"{pad}    ROW: " + " || ".join(c.text for c in row.cells))

        if sh.shape_type == 13:  # PICTURE
            try:
                print(f"{pad}    PICTURE: {sh.image.filename or sh.image.content_type}")
            except ValueError:
                print(f"{pad}    PICTURE: (no embedded raster; likely SVG)")

        if getattr(sh, "has_chart", False) and sh.has_chart:
            ch = sh.chart
            print(f"{pad}    CHART: {ch.chart_type}")
            try:
                print(f"{pad}      cats: {list(ch.plots[0].categories)}")
            except Exception:
                pass
            for s in ch.plots[0].series:
                print(f"{pad}      series {s.name!r}: {list(s.values)}")


def main():
    path = sys.argv[1]
    wanted = {int(a) for a in sys.argv[2:]}
    prs = Presentation(path)
    print(f"{path}: {len(prs.slides)} slides, {prs.slide_width}x{prs.slide_height}")
    print("layouts: " + ", ".join(f"{i}:{l.name}" for i, l in enumerate(prs.slide_layouts)))
    for n, slide in enumerate(prs.slides, 1):
        if wanted and n not in wanted:
            continue
        print("\n" + "=" * 72)
        print(f"SLIDE {n}  (layout: {slide.slide_layout.name})")
        print("=" * 72)
        walk(slide.shapes)
        if slide.has_notes_slide:
            tf = slide.notes_slide.notes_text_frame
            if tf and tf.text.strip():
                print("--- NOTES ---")
                print(tf.text)


if __name__ == "__main__":
    main()
