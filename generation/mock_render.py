#!/usr/bin/env python3
"""Approximate-render a pptx to PNG via HTML + headless Chrome.

There is no LibreOffice on this machine, so this is the only way to *see* a deck.
It reads real geometry, fills, text runs, and embedded images out of the file, so
it catches overflow, collisions, and bad spacing. It is NOT a faithful renderer:
fonts fall back to a generic sans, autofit and bullet glyphs are ignored, and
charts are drawn as labelled grey boxes.

    python3 generation/mock_render.py "slides/Valkey placeholder slides.pptx" [out.png]
"""
import os
import shutil
import subprocess
import sys
import tempfile

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
SCALE = 10160.0          # EMU per px -> 1200px-wide slide
PT_PX = 12700 / SCALE    # px per point

# Valkey theme defaults, keyed by both the XML scheme names (from shape fills) AND
# the python-pptx enum member names (from run.font.color.theme_color), because those
# two vocabularies differ.
#
# load_theme() overwrites these from the deck's own clrScheme. Hardcoding them was a
# bug with teeth: build_placeholder_deck.py patches accent6 from the stock peach to
# the brand rose, and this table kept rendering the old peach, so a correct deck
# looked wrong. Read the file, do not assume the theme.
THEME = {
    # XML scheme colour names
    "dk1": "#000000", "lt1": "#FFFFFF", "dk2": "#232F3E", "lt2": "#F1F3F3",
    "accent1": "#6983FF", "accent2": "#1A2026", "accent3": "#632637",
    "accent4": "#30176E", "accent5": "#BCB5E7", "accent6": "#FBD8BF",
    "tx1": "#000000", "bg1": "#FFFFFF", "tx2": "#232F3E", "bg2": "#F1F3F3",
    "hlink": "#6983FF", "folHlink": "#6983FF",
    # MSO_THEME_COLOR enum member names
    "dark_1": "#000000", "dark_2": "#232F3E",
    "light_1": "#FFFFFF", "light_2": "#F1F3F3",
    "text_1": "#000000", "text_2": "#232F3E",
    "background_1": "#FFFFFF", "background_2": "#F1F3F3",
    "accent_1": "#6983FF", "accent_2": "#1A2026", "accent_3": "#632637",
    "accent_4": "#30176E", "accent_5": "#BCB5E7", "accent_6": "#FBD8BF",
    "hyperlink": "#6983FF", "followed_hyperlink": "#6983FF",
}
UNKNOWN_COLORS = set()

# python-pptx enum names, not the XML "ctr"/"r" values.
ALIGN = {"center": "center", "right": "right", "justify": "justify", "left": "left"}


def px(emu):
    return (emu or 0) / SCALE


def scheme_hex(el):
    """Resolve a solidFill element to a hex color."""
    if el is None:
        return None
    sc = el.find(qn("a:schemeClr"))
    if sc is not None:
        return THEME.get(sc.get("val"), "#888888")
    srgb = el.find(qn("a:srgbClr"))
    if srgb is not None:
        return "#" + srgb.get("val")
    return None


def grad_css(gradFill):
    """A linear a:gradFill as a CSS linear-gradient.

    Worth supporting rather than ignoring: a gradient with a hard stop is how a
    one-sided accent edge is done in a single shape, and skipping gradients let a
    malformed one (two stops at the same position, which renders as a full-width wash)
    ship unnoticed. Positions are 1/1000 of a percent. a:lin/@ang is 60000ths of a
    degree clockwise from "to the right", and CSS 0deg is "to the top", so the two
    conventions differ by 90.
    """
    gsLst = gradFill.find(qn("a:gsLst"))
    if gsLst is None:
        return None
    stops = []
    for gs in gsLst.findall(qn("a:gs")):
        colour = scheme_hex(gs)
        if colour is None:
            return None
        stops.append(f"{colour} {int(gs.get('pos', 0)) / 1000:.3f}%")
    if len(stops) < 2:
        return None
    lin = gradFill.find(qn("a:lin"))
    deg = (int(lin.get("ang", 0)) / 60000 if lin is not None else 0) + 90
    return f"linear-gradient({deg:g}deg, {', '.join(stops)})"


def shape_fill(sp):
    """CSS `background` value for a shape, or None. May be a colour or a gradient."""
    spPr = sp.find(qn("p:spPr"))
    if spPr is None:
        return None
    if spPr.find(qn("a:noFill")) is not None:
        return None
    solid = scheme_hex(spPr.find(qn("a:solidFill")))
    if solid:
        return solid
    gradFill = spPr.find(qn("a:gradFill"))
    return grad_css(gradFill) if gradFill is not None else None


def placeholder_fill(shape, layout):
    """A slide placeholder's fill, falling back to its layout placeholder.

    Placeholders inherit shape properties, so a row whose card background lives on the
    layout has an empty spPr of its own. Without this the mock draws nothing for it and
    the slide looks bare when it is not.
    """
    own = shape_fill(shape._element)
    if own or layout is None:
        return own
    try:
        idx = shape.placeholder_format.idx
    except (ValueError, AttributeError):
        return None
    for ph in layout.placeholders:
        if ph.placeholder_format.idx == idx:
            return shape_fill(ph._element)
    return None


def shape_line(sp):
    spPr = sp.find(qn("p:spPr"))
    ln = spPr.find(qn("a:ln")) if spPr is not None else None
    if ln is None or ln.find(qn("a:noFill")) is not None:
        return None
    c = scheme_hex(ln.find(qn("a:solidFill")))
    # `w` is optional: absent means the DrawingML default of 0 EMU, a hairline.
    w = int(ln.get("w") or 0)
    return (c, max(1.0, w / 12700 * PT_PX)) if c else None


def run_style(run, default_pt, default_color, default_bold):
    f = run.font
    size = f.size.pt if f.size is not None else default_pt
    bold = default_bold if f.bold is None else f.bold
    color = default_color
    try:
        if f.color is not None and f.color.type is not None:
            if f.color.type == 2:  # theme color
                key = str(f.color.theme_color).split()[0].lower()
                if key in THEME:
                    color = THEME[key]
                else:
                    # Do not silently fall back: an unmapped colour renders as the
                    # default and quietly misrepresents the slide.
                    UNKNOWN_COLORS.add(key)
                    color = "#ff00ff"
            else:
                color = "#" + str(f.color.rgb)
    except Exception:
        pass
    return size, bold, color, f.italic


def text_html(shape, is_title):
    if not shape.has_text_frame:
        return ""
    d_pt = 32 if is_title else 18
    d_bold = bool(is_title)
    out = []
    for p in shape.text_frame.paragraphs:
        if not p.runs:
            out.append("<div style='height:0.4em'></div>")
            continue
        al = ALIGN.get(str(p.alignment).split()[0].lower() if p.alignment else "", "left")
        ls = f"line-height:{p.line_spacing};" if isinstance(p.line_spacing, float) else "line-height:1.25;"
        spans = []
        for r in p.runs:
            size, bold, color, ital = run_style(r, d_pt, "#000000", d_bold)
            spans.append(
                f"<span style='font-size:{size * PT_PX:.1f}px;"
                f"font-weight:{700 if bold else 400};color:{color};"
                f"{'font-style:italic;' if ital else ''}'>"
                f"{(r.text or '').replace('&','&amp;').replace('<','&lt;')}</span>")
        out.append(f"<div style='text-align:{al};{ls}'>{''.join(spans)}</div>")
    return "".join(out)


SVG_NS = "{http://schemas.microsoft.com/office/drawing/2016/SVG/main}"


def svg_or_png(shape, slide):
    """Return (blob, ext), preferring the SVG part when the pic carries an svgBlip.

    `shape.image` resolves the a:blip r:embed, which is the PNG fallback. PowerPoint
    actually renders the svgBlip, so the mock should too.
    """
    blipFill = shape._element.find(qn("p:blipFill"))
    blip = blipFill.find(qn("a:blip")) if blipFill is not None else None
    if blip is not None:
        for ext_el in blip.iterfind(f"{qn('a:extLst')}/{qn('a:ext')}"):
            sb = ext_el.find(f"{SVG_NS}svgBlip")
            if sb is not None:
                rId = sb.get(qn("r:embed"))
                return slide.part.related_part(rId).blob, "svg"
    img = shape.image
    return img.blob, (img.ext or "png")


# clrScheme name -> every THEME key that resolves to it. dk1/lt1 are mapped to
# tx1/bg1 by the master's clrMap, which is the default mapping in every deck here.
_SCHEME_KEYS = {
    "dk1": ("dk1", "tx1", "dark_1", "text_1"),
    "lt1": ("lt1", "bg1", "light_1", "background_1"),
    "dk2": ("dk2", "tx2", "dark_2", "text_2"),
    "lt2": ("lt2", "bg2", "light_2", "background_2"),
    "hlink": ("hlink", "hyperlink"),
    "folHlink": ("folHlink", "followed_hyperlink"),
    **{f"accent{i}": (f"accent{i}", f"accent_{i}") for i in range(1, 7)},
}


def load_theme(prs):
    """Overwrite THEME from the deck's own colour scheme."""
    for part in prs.part.package.iter_parts():
        if not str(part.partname).startswith("/ppt/theme/theme1.xml"):
            continue
        root = etree.fromstring(part.blob)
        scheme = root.find(f".//{qn('a:clrScheme')}")
        if scheme is None:
            return
        for entry in scheme:
            srgb = entry.find(qn("a:srgbClr"))
            name = etree.QName(entry).localname
            if srgb is None or name not in _SCHEME_KEYS:
                continue
            for key in _SCHEME_KEYS[name]:
                THEME[key] = "#" + srgb.get("val")
        return


def render(path, out_png):
    prs = Presentation(path)
    load_theme(prs)
    tmp = tempfile.mkdtemp(prefix="mockrender-")
    W, H = prs.slide_width, prs.slide_height
    blocks = []

    for idx, slide in enumerate(prs.slides, 1):
        parts = [f"<div class=s style='width:{px(W)}px;height:{px(H)}px'>"
                 f"<div class=n>{idx}</div>"]
        title_el = slide.shapes.title._element if slide.shapes.title is not None else None

        # Master then layout content underneath the slide's own shapes. These carry
        # the corner Valkey logo and the dark full-bleed backgrounds on the title
        # layouts, so omitting them misrepresents the slide badly.
        layout = slide.slide_layout
        inherited = []
        if slide._element.get("showMasterSp") != "0":
            inherited += list(layout.slide_master.shapes)
        inherited += list(layout.shapes)
        for shape in inherited:
            if shape.is_placeholder:
                continue          # layout placeholders hold "Click to add..." prompts
            if shape.left is None:
                continue
            box = (f"position:absolute;left:{px(shape.left):.1f}px;top:{px(shape.top):.1f}px;"
                   f"width:{px(shape.width):.1f}px;height:{px(shape.height):.1f}px;")
            if shape.shape_type == 13:
                try:
                    img = shape.image
                    fn = os.path.join(tmp, f"m{idx}_{abs(hash(box))}.{img.ext or 'png'}")
                    with open(fn, "wb") as fh:
                        fh.write(img.blob)
                    parts.append(f"<img src='file://{fn}' style='{box}'>")
                except Exception:
                    pass
                continue
            fill = shape_fill(shape._element)
            if fill:
                parts.append(f"<div style='{box}background:{fill}'></div>")

        for shape in slide.shapes:
            el = shape._element
            L, T, Wd, Ht = px(shape.left), px(shape.top), px(shape.width), px(shape.height)
            box = f"position:absolute;left:{L:.1f}px;top:{T:.1f}px;width:{Wd:.1f}px;height:{Ht:.1f}px;"

            if shape.shape_type == 13 or el.tag == qn("p:pic"):
                try:
                    blob, ext = svg_or_png(shape, slide)
                    fn = os.path.join(tmp, f"i{idx}_{abs(hash(str(L) + str(T) + ext))}.{ext}")
                    with open(fn, "wb") as fh:
                        fh.write(blob)
                    parts.append(f"<img src='file://{fn}' style='{box}'>")
                except Exception as exc:
                    parts.append(f"<div style='{box}outline:1px dashed #f0f;font:9px sans-serif;"
                                 f"color:#f0f'>img: {exc}</div>")
                continue

            if getattr(shape, "has_chart", False) and shape.has_chart:
                parts.append(f"<div style='{box}background:#e9e9e9;border:1px solid #bbb;"
                             f"font:12px sans-serif;color:#666;display:flex;align-items:center;"
                             f"justify-content:center'>chart: {shape.chart.chart_type}</div>")
                continue

            if el.tag == qn("p:cxnSp"):
                # A connector is a line across its bounding box, not a box.
                xfrm = el.find(".//" + qn("a:xfrm"))
                fh, fv = xfrm.get("flipH") == "1", xfrm.get("flipV") == "1"
                geom = el.find(".//" + qn("a:prstGeom"))
                prst = geom.get("prst") if geom is not None else "line"
                # Path in unflipped box coordinates, then mirror for flipH / flipV.
                if prst == "bentConnector2":
                    pts = [(0, 0), (Wd, 0), (Wd, Ht)]
                elif prst == "bentConnector3":
                    pts = [(0, 0), (Wd / 2, 0), (Wd / 2, Ht), (Wd, Ht)]
                else:
                    pts = [(0, 0), (Wd, Ht)]
                pts = [((Wd - x) if fh else x, (Ht - y) if fv else y) for x, y in pts]
                colour, width = shape_line(el) or ("#000", 1)
                dash = "stroke-dasharray='6 4'" if el.find(".//" + qn("a:prstDash")) is not None else ""
                path = " ".join(f"{x:.1f},{y:.1f}" for x, y in pts)
                parts.append(f"<svg style='position:absolute;left:{L:.1f}px;top:{T:.1f}px;"
                             f"overflow:visible' width='{max(Wd, 1):.1f}' height='{max(Ht, 1):.1f}'>"
                             f"<polyline points='{path}' fill='none' "
                             f"stroke='{colour}' stroke-width='{max(width, 1):.1f}' {dash}/></svg>")
                continue

            fill = (placeholder_fill(shape, layout) if shape.is_placeholder
                    else shape_fill(el))
            line = shape_line(el)
            # border-box always: PowerPoint's text margins are internal and do not
            # change a shape's size, but as CSS padding they inflate a content-box
            # div. Without this every filled shape renders wider than it really is.
            style = box + "box-sizing:border-box;"
            if fill:
                style += f"background:{fill};"
            if line:
                style += f"border:{line[1]:.1f}px solid {line[0]};"
            inner = text_html(shape, el is title_el)
            pad = ""
            if shape.has_text_frame:
                tf = shape.text_frame
                # python-pptx enum names (MIDDLE, BOTTOM), not the XML values
                anchor = {"middle": "center", "bottom": "flex-end"}.get(
                    str(tf.vertical_anchor).split()[0].lower() if tf.vertical_anchor else "", "flex-start")
                pad = (f"display:flex;flex-direction:column;justify-content:{anchor};"
                       f"padding:{px(tf.margin_top or 0):.1f}px {px(tf.margin_right or 0):.1f}px "
                       f"{px(tf.margin_bottom or 0):.1f}px {px(tf.margin_left or 0):.1f}px;")
            parts.append(f"<div style='{style}{pad}font-family:sans-serif'>{inner}</div>")

        parts.append("</div>")
        blocks.append("".join(parts))

    html = ("<html><body style='margin:0;background:#777'>"
            "<style>.s{position:relative;background:#fff;margin:0 auto 14px;overflow:hidden}"
            ".n{position:absolute;right:2px;top:1px;font:10px sans-serif;color:#bbb;z-index:99}"
            "</style>" + "".join(blocks) + "</body></html>")
    html_path = os.path.join(tmp, "deck.html")
    with open(html_path, "w") as fh:
        fh.write(html)

    height = int(px(H) + 14) * len(prs.slides) + 20
    subprocess.run(
        [CHROME, "--headless", "--disable-gpu", "--no-sandbox",
         f"--screenshot={out_png}", f"--window-size={int(px(W)) + 15},{height}",
         "--force-device-scale-factor=1.4", f"file://{html_path}"],
        check=True, capture_output=True)
    shutil.rmtree(tmp, ignore_errors=True)
    print(f"{out_png}  ({len(prs.slides)} slides)")
    if UNKNOWN_COLORS:
        print(f"WARNING unmapped theme colours (drawn magenta): {sorted(UNKNOWN_COLORS)}")


if __name__ == "__main__":
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else "/tmp/deck_mock.png"
    render(src, dst)
