#!/usr/bin/env python3
"""Build "Open-Source Maintainers vs. The CVE Tsunami" from the placeholder deck.

    python3 generation/build_cve_tsunami_deck.py [out.pptx]

The spec is the vault outline, `content/talks/Open Source Maintainers Vs. The CVE
Tsunami.md`: every title, every piece of on-screen text, and every note comes from it.
Change the outline first, then this script. Never hand-edit the generated pptx.

Speakers: Madelyn Olson and Harkrishn (Hari) Patro.

The donor is the placeholder deck, because its master carries the twelve Valkey layouts.
Diagrams are native shapes. Build frames are separate slides drawn by the same function
with an increasing frame number, so each frame only ever adds.

Code is set in Courier New, the only exception to the Arial rule: a code listing in a
proportional font loses its indentation, and Courier New is on every machine Arial is.
"""
import os
import shutil
import sys

from pptx import Presentation
from pptx.enum.dml import MSO_LINE, MSO_THEME_COLOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_placeholder_deck as V
from build_placeholder_deck import (CARD_ICON, CARD_ROSE, CONTENT_LEFT, PAD,
                                    SLIDE_H, SLIDE_W, arrow, box_columns,
                                    card_geometry, rect, source_note, text_box)
from build_nerdearla_deck import (audit_list_rows, audit_paragraph_order, drop_placeholders,
                                  fill, list_rows, notes, on_layout, retypeface,
                                  set_title, shapes_flat)
from pptx_svg import add_svg

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)                      # repo root
SLIDES = os.path.join(ROOT, "slides")             # generated decks live here
TEMPLATE = os.path.join(SLIDES, "Valkey placeholder slides.pptx")
OUT = os.path.join(SLIDES, "Open Source Maintainers vs The CVE Tsunami.pptx")
ASSETS = os.path.join(HERE, "assets")

FONT, CODE_FONT = "Arial", "Courier New"
BLOG = "Source: valkey.io/blog/keeping-up-with-ai-valkey-security"

AI = MSO_THEME_COLOR.ACCENT_1          # AI work: Valkey blue, white text
HUMAN = CARD_ROSE                      # human work: brand rose, dark text
DARK = MSO_THEME_COLOR.ACCENT_2
GREY = MSO_THEME_COLOR.LIGHT_2
WHITE = MSO_THEME_COLOR.LIGHT_1
INK = MSO_THEME_COLOR.DARK_1
INK2 = MSO_THEME_COLOR.DARK_2
LILAC = MSO_THEME_COLOR.ACCENT_5

GRID_L = CONTENT_LEFT
GRID_W = SLIDE_W - 2 * CONTENT_LEFT    # symmetric, see slides-skill.md Standard Geometry
MID_Y = SLIDE_H // 2


# --- primitives ---------------------------------------------------------------


def icon(slide, name, x, y, side, on=None):
    """An icon. Pass `on=` the fill it sits on: rose gets the black variant."""
    add_svg(slide, *V.icon_paths(name, on), int(x), int(y), int(side), int(side),
            name=name)


def label(slide, x, y, w, h, text, size=18, bold=True, colour=INK, align=PP_ALIGN.CENTER):
    return text_box(slide, x, y, w, h, [(text, size, bold, colour)], align=align,
                    anchor=MSO_ANCHOR.MIDDLE)


def block(slide, x, y, w, h, text, fill, fg, size=18, line=None, dashed=False):
    """A filled box whose label names it. The label is on the shape, never under it.

    Plain boxes are the template's `chip()`. Only the outlined and dashed variants
    (the freed connection, the attacker bytes) are drawn here.
    """
    if line is None and not dashed and fill is not None:
        V.chip(slide, x, y, w, h, text, fill=fill, fg=fg, size=size)
        return
    sh = rect(slide, x, y, w, h, fill=fill, line=line)
    if dashed:
        sh.line.dash_style = MSO_LINE.DASH
        sh.line.width = Pt(2)
    label(slide, x, y, w, h, text, size=size, colour=fg)
    return sh


def code_block(slide, x, y, w, lines, size=18, hl=()):
    """Monospace listing on a dark panel. `hl` holds 0-based lines to highlight."""
    line_h = Pt(size * 1.2)
    h = len(lines) * line_h + 2 * PAD
    rect(slide, x, y, w, h, fill=DARK)
    box = slide.shapes.add_textbox(Emu(int(x + PAD)), Emu(int(y + PAD)),
                                   Emu(int(w - 2 * PAD)), Emu(int(h - 2 * PAD)))
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = para.add_run()
        r.text = line
        r.font.name = CODE_FONT
        r.font.size = Pt(size)
        r.font.bold = i in hl
        r.font.color.theme_color = HUMAN if i in hl else WHITE
    return h


def card_lines(slide, x, y, w, h, fill, icon_name, heading, lines, text=V.CARD_TEXT,
               bold_first=False):
    """card(), with one paragraph per line in the body instead of one sentence."""
    if fill == HUMAN:
        text = INK                      # black text on rose, always
    rect(slide, x, y, w, h, fill=fill)
    icon_r, head_r, body_r = card_geometry(x, y, w, h)
    icon(slide, icon_name, *icon_r[:2], icon_r[2], on=fill)
    text_box(slide, *head_r, [(heading, 18, True, text)])
    if not lines:
        return
    box = slide.shapes.add_textbox(*(Emu(int(v)) for v in body_r))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.space_after = Pt(6)
        r = para.add_run()
        r.text = line
        r.font.name = FONT
        r.font.size = Pt(16)
        r.font.bold = bold_first and i == 0
        r.font.color.theme_color = text


def cards(prs, title, specs, fills, h, bold_first=False):
    """n cards, centred on the page, on `Valkey Title Only` (custom icons)."""
    slide = on_layout(prs, V.LAYOUT_TITLE_ONLY)
    set_title(slide, title)
    col_w, xs = box_columns(len(specs))
    top = (SLIDE_H - h) // 2 + 200000
    for x, f, spec in zip(xs, fills, specs):
        card_lines(slide, x, top, col_w, h, f, *spec, bold_first=bold_first)
    return slide


def titled(prs, title):
    slide = on_layout(prs, V.LAYOUT_TITLE_ONLY)
    set_title(slide, title)
    return slide


def statement(prs, text, caption=None, bold=False, size=None):
    slide = on_layout(prs, V.LAYOUT_INSIGHT_BOLD if bold else V.LAYOUT_INSIGHT)
    fill(slide, 1, text, size=size)
    if caption:
        fill(slide, 2, caption)
    else:
        drop_placeholders(slide, 2)
    return slide


def big_number(prs, number, text, source):
    slide = on_layout(prs, V.LAYOUT_BIG_NUMBER)
    fill(slide, 1, number)
    fill(slide, 2, text)
    fill(slide, 3, source)
    return slide


# --- diagram slides ------------------------------------------------------------


# Published advisories in valkey-io/valkey, by month, from
# `gh api repos/valkey-io/valkey/security-advisories?state=published` on 2026-10-07.
# None is a month that doesn't count: before the fork (Jan, Feb 2024) or not yet
# happened (Nov, Dec 2026). October 2026 is partial.
ADVISORIES = {
    "2024": [None, None, 0, 0, 0, 0, 0, 0, 0, 0, 3, 0],
    "2025": [0, 0, 0, 1, 0, 0, 0, 0, 0, 1, 0, 0],
    "2026": [0, 3, 0, 0, 0, 0, 2, 2, 0, 0, None, None],
}
ADVISORY_SOURCE = "Source: github.com/valkey-io/valkey/security/advisories"


def slide_advisories(prs):
    """Three bars, one per year: published advisories, absolute counts."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
    slide = titled(prs, "2026 already has more advisories than 2024 and 2025 combined")
    counts = {y: sum(v for v in vals if v) for y, vals in ADVISORIES.items()}
    cd = CategoryChartData()
    cd.categories = list(counts)
    cd.add_series("advisories", list(counts.values()))
    left, top, w = GRID_L + 1200000, 1350000, GRID_W - 2400000
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(left), Emu(top),
                                   Emu(w), Emu(5850000 - top), cd).chart
    chart.font.name = FONT
    chart.font.size = Pt(20)
    chart.has_legend = False
    plot = chart.plots[0]
    plot.gap_width = 80
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = "0", False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.bold = Pt(28), True
    for point, year in zip(plot.series[0].points, counts):
        point.format.fill.solid()
        point.format.fill.fore_color.theme_color = AI if year == "2026" else LILAC
    va = chart.value_axis
    va.minimum_scale = 0
    va.visible = False
    va.has_major_gridlines = False
    chart.category_axis.format.line.color.theme_color = INK2
    source_note(slide, ADVISORY_SOURCE)
    notes(slide, "And it’s not just that. We’ve observed a increase in advisory in Valkey this year. GitHub published more security advisories this year than the combined advisories from the last two years.")
    return slide


def slide_three_reporters(prs):
    slide = titled(prs, "Finding bugs got easy")
    side, pitch, x = 900000, 1200000, 1700000
    top = MID_Y - (pitch * 2 + side) // 2 + 150000
    tx, tw, th = 6900000, 4300000, 1000000
    ty = top + pitch + side // 2 - th // 2
    block(slide, tx, ty, tw, th, "script debugger use-after-free", DARK, WHITE)
    for i in range(3):
        y = top + i * pitch
        icon(slide, "report", x, y, side)
        arrow(slide, x + side + 150000, y + side // 2, tx - 100000, ty + th // 2,
              colour=INK2, elbow=None if i == 1 else 3)
    label(slide, x - 300000, top + 2 * pitch + side + 60000, side + 600000, 380000,
          "2 weeks", size=16, bold=False, colour=INK2)
    notes(slide, "Three researchers who didn't know about each other reported the same use-after-free in the script debugger, all inside two weeks. Finding bugs is real cheap now. What follows next is expensive.")
    return slide


def slide_pipeline(prs):
    slide = titled(prs, "Finding got cheap. Everything after it didn't.")
    steps = ["find", "reproduce", "assess", "fix", "embargo", "backport", "release"]
    gap, first_w, h = 80000, 760000, 1700000
    rest_w = (GRID_W - first_w - gap * (len(steps) - 1)) // (len(steps) - 1)
    y, x = MID_Y - h // 2 + 150000, GRID_L
    for i, step in enumerate(steps):
        w = first_w if i == 0 else rest_w
        # No icon on the blue box: the icons are Valkey-blue strokes and vanish on it.
        bg, fg, ic = (AI, WHITE, None) if i == 0 else (HUMAN, INK, "person")
        rect(slide, x, y, w, h, fill=bg)
        if ic:
            icon(slide, ic, x + (w - 560000) // 2, y + 280000, 560000, on=bg)
        label(slide, x, y + 1000000, w, 420000, step, size=16, colour=fg)
        x += w + gap
    notes(slide, "Finding is the only step a model made cheap. Every step after it is "
                 "maintainer judgement or coordination, and none of them scale with "
                 "the number of reports.")
    return slide


def funnel(slide, stages, top, row_h=700000, gap=150000):
    """Centred, narrowing rows. `stages` is [(text, width, fill, fg, icon or None)]."""
    for i, (text, w, f, fg, ic) in enumerate(stages):
        x, y = (SLIDE_W - w) // 2, top + i * (row_h + gap)
        rect(slide, x, y, w, row_h, fill=f)
        if ic:
            icon(slide, ic, x + 150000, y + (row_h - 480000) // 2, 480000, on=f)
        label(slide, x, y, w, row_h, text, colour=fg)
    return top + len(stages) * (row_h + gap) - gap


def slide_audit_funnel(prs, frame):
    slide = titled(prs, "Humans only see crashing tests")
    stages = [("threat model", 9600000, AI, WHITE, None),
              ("AI threat generation", 8400000, AI, WHITE, None),
              ("AI review", 7200000, AI, WHITE, None),
              ("reproduce in harness", 3600000, LILAC, DARK, None)][:frame]
    top = 1250000
    bottom = funnel(slide, stages, top)
    if frame == 4:
        icon(slide, "person", (SLIDE_W - 620000) // 2, bottom + 180000, 620000)
    notes(slide, "AI threat generation proposes candidate bugs, then AI review argues against each. Most "
                 "die at verification: the first model misread the code, or the "
                 "impact is trivial. The harness is deterministic: crash, ASAN failure, "
                 "replica desync. If the harness can't tell, the model can't prove it.")
    return slide


NODE_W, NODE_H, NODE_Y = 1700000, 1000000, 2900000
NODE_X = {"A": 1500000, "B": 5246000, "C": 8992000}


def slide_tls_list(prs, frame):
    slide = titled(prs, "The TLS use-after-free")
    ax, bx, cx = NODE_X["A"], NODE_X["B"], NODE_X["C"]
    mid = NODE_Y + NODE_H // 2
    # Connections, not Valkey instances, so no corner glyph.
    V.node(slide, ax, NODE_Y, NODE_W, NODE_H, "A", mark=False)
    if frame >= 3:
        block(slide, bx, NODE_Y, NODE_W, NODE_H, "B", None, INK2, size=16,
              line=INK2, dashed=True)
    else:
        V.node(slide, bx, NODE_Y, NODE_W, NODE_H, "B", mark=False)
    V.node(slide, cx, NODE_Y, NODE_W, NODE_H, "C", mark=False)
    arrow(slide, ax + NODE_W, mid, bx, mid, colour=INK2)
    arrow(slide, bx + NODE_W, mid, cx, mid, colour=INK2)

    chip_h, chip_y = 480000, NODE_Y - 1150000
    block(slide, ax, chip_y, NODE_W, chip_h, "iterator", DARK, WHITE, size=16)
    arrow(slide, ax + NODE_W // 2, chip_y + chip_h, ax + NODE_W // 2, NODE_Y, colour=DARK)
    nxt = V.FAIL_RED if frame >= 4 else AI
    arrow(slide, ax + NODE_W, chip_y + chip_h // 2, bx + NODE_W // 2, NODE_Y,
          colour=nxt, pt=2.25, dashed=frame < 4, elbow=2)
    label(slide, ax + NODE_W + 250000, chip_y - 380000, 1400000, 360000, "next",
          size=16, colour=nxt)

    if frame >= 2:
        ky = NODE_Y + NODE_H + 700000
        block(slide, ax - 150000, ky, NODE_W + 300000, chip_h, "CLIENT KILL B", HUMAN,
              DARK, size=16)
        arrow(slide, ax + NODE_W + 150000, ky + chip_h // 2, bx + NODE_W // 2,
              NODE_Y + NODE_H, colour=INK2, elbow=2)
    if frame >= 3:
        # Beside the CLIENT KILL arrow's vertical leg, not on top of it.
        label(slide, bx + NODE_W // 2 + 100000, NODE_Y + NODE_H + 80000, 1000000,
              380000, "freed", size=16, colour=INK2, align=PP_ALIGN.LEFT)
    if frame >= 4:
        side = 700000
        icon(slide, "crash", bx + NODE_W + 120000, chip_y - 250000, side)
        label(slide, bx + NODE_W + 120000 + side, chip_y - 250000, 1300000, side,
              "crash", size=20, colour=V.FAIL_RED, align=PP_ALIGN.LEFT)
    notes(slide, "TLS can decrypt more than one command's worth of data, so connections "
                 "with unread data go on a list walked once per event loop pass. The "
                 "walk holds a pointer to the next connection.\n\n"
                 "The freed object is one the server calls through, so attacker bytes "
                 "there plausibly mean code execution. That's why it's a CVE under "
                 "Act 1's rule.")
    return slide


TLS_C = [
    "listRewind(pending_list, &li);",
    "while ((ln = listNext(&li))) {",
    "    tls_connection *conn = listNodeValue(ln);",
    "    if (conn->flags & TLS_CONN_FLAG_POSTPONE_UPDATE_STATE) continue;",
    "    tlsHandleEvent(conn, AE_READABLE);",
    "    processed++;",
    "}",
]


def slide_tls_code(prs):
    slide = titled(prs, "The TLS use-after-free")
    h = len(TLS_C) * Pt(18 * 1.2) + 2 * PAD
    code_block(slide, GRID_L, MID_Y - h // 2 + 150000, GRID_W, TLS_C, hl=(1,))
    source_note(slide, "src/tls.c, tlsProcessPendingData(), before the fix")
    notes(slide, "listIter caches the next node. tlsHandleEvent can run CLIENT KILL -> "
                 "freeClient -> connTLSClose -> listDelNode, freeing that cached node. "
                 "The next listNext() reads freed memory.\n\n"
                 "The fix pops each connection off the head before handling it, bounded "
                 "by the length captured on entry. Commit \"Fix use-after-free in "
                 "tlsProcessPendingData() on CLIENT KILL (#4234)\", GHSA-53mc-f3m3-99vh.")
    return slide


def slide_what_about_rust(prs):
    """The question the room is already asking, with Ferris, before any Rust code."""
    slide = statement(prs, "What about Rust?")
    path = os.path.join(ASSETS, "ferris.png")       # CC0, rustacean.net
    from PIL import Image
    w_px, h_px = Image.open(path).size
    w = 3600000
    h = int(w * h_px / w_px)
    slide.shapes.add_picture(path, Emu(GRID_L + GRID_W - w), Emu(MID_Y - h // 2),
                             Emu(w), Emu(h))
    notes(slide, "Someone is going to ask. Memory-safe languages are the obvious answer "
                 "to a use-after-free, so take the question head on.")
    return slide


RUST = [
    "for conn in server.tls_pending.iter() {",
    "    if conn.postpone { continue; }",
    "    server.handle_event(conn); // CLIENT KILL removes from tls_pending",
    "}",
]
RUSTC = [
    "error[E0502]: cannot borrow `*server` as mutable",
    "              because it is also borrowed as immutable",
]


def slide_rust(prs, frame):
    slide = titled(prs, "Rust won't compile it")
    line_h = Pt(18 * 1.2)
    h1 = len(RUST) * line_h + 2 * PAD
    h2 = len(RUSTC) * line_h + 2 * PAD
    top = MID_Y - (h1 + 250000 + h2) // 2 + 150000
    code_block(slide, GRID_L, top, GRID_W, RUST, hl=(2,) if frame == 2 else ())
    if frame == 2:
        code_block(slide, GRID_L, top + h1 + 250000, GRID_W, RUSTC, hl=(0, 1))
    notes(slide, "Same loop as the C. handle_event needs &mut server because it can "
                 "remove from the list, and the iterator already holds a borrow of it. "
                 "Checked with rustc --edition 2021; the error is real output, wrapped "
                 "to fit.\n\n"
                 "The version that compiles is the shape of the C fix: pop_front() each "
                 "connection, bounded by the length on entry. The compiler forces the "
                 "fix that took us a CVE to find.")
    return slide


def slide_rust_rows(prs):
    """What Rust removes in black, what it leaves in red."""
    slide = on_layout(prs, V.LAYOUT_LIST_ROWS)
    set_title(slide, "Rust removes some bugs")
    removed = ["Use after free", "Access out of bounds"]
    kept = ["Panics", "Hangs", "Logic bugs"]
    list_rows(slide, removed + kept)
    for ph in slide.placeholders:
        idx = ph.placeholder_format.idx
        if not 1 <= idx <= len(removed) + len(kept):
            continue
        for run in ph.text_frame.paragraphs[0].runs:
            # Red is the one explicit-RGB exception to theme colours (failure states).
            V._set_colour(run.font.color, INK if idx <= len(removed) else V.FAIL_RED)
    notes(slide, "The TLS bug was a use after free, and safe Rust won't compile it. "
                 "Out-of-bounds access is checked at runtime: it becomes a panic, so no "
                 "leaked bytes, but still a crash. Panics, hangs, and a permission check in "
                 "the wrong place all still ship.")
    return slide

def slide_acl(prs):
    slide = titled(prs, "Doing what you're allowed isn't a vulnerability")
    line_y = MID_Y + 150000
    rect(slide, GRID_L, line_y, GRID_W, 45720, fill=DARK)
    label(slide, GRID_L + GRID_W - 2500000, line_y - 420000, 2500000, 380000,
          "ACL boundary", size=14, bold=False, colour=INK2, align=PP_ALIGN.RIGHT)
    gap, h = 200000, 750000
    above = ["read another client's keys", "execute code"]
    below = ["delete data", "grow memory", "pin a core with Lua"]
    w = (GRID_W - 2 * gap) // 3
    for i, t in enumerate(above):
        block(slide, GRID_L + i * (w + gap), line_y - 600000 - h, w, h, t, HUMAN, INK)
    for i, t in enumerate(below):
        block(slide, GRID_L + i * (w + gap), line_y + 600000, w, h, t, GREY, INK2)
    notes(slide, "Below the line is what an authenticated client can already do. Above it is what crosses a permission boundary. Out-of-bounds reads cross the line because the leaked bytes can belong to another client. We’re raising the bar for what counts as a CVE.")
    return slide


def slide_cluster_bus(prs):
    """Cluster networking: clients on the client port, nodes on the bus, behind a fence."""
    slide = titled(prs, "The cluster bus trusts every message")
    nw, nh, ny = 2200000, 900000, 3150000
    xs = [1500000, 5000000, 8500000]
    cx = [x + nw // 2 for x in xs]
    cw, ch, cy = 3000000, 500000, 1300000
    V.chip(slide, (SLIDE_W - cw) // 2, cy, cw, ch, "clients", fill=DARK, fg=WHITE, size=16)
    bar_y = 2350000
    arrow(slide, SLIDE_W // 2, cy + ch, SLIDE_W // 2, bar_y, colour=INK2, head=False)
    arrow(slide, cx[0], bar_y, cx[2], bar_y, colour=INK2, head=False)
    for c in cx:
        arrow(slide, c, bar_y, c, ny, colour=INK2)
        label(slide, c + 60000, bar_y + 80000, 900000, 340000, "6379", size=14,
              bold=False, colour=INK2, align=PP_ALIGN.LEFT)
    fence = rect(slide, 1150000, 2900000, 9900000, 2300000, line=INK2)
    fence.line.dash_style = MSO_LINE.DASH
    label(slide, 1250000, 4800000, 3000000, 340000, "nodes only", size=14, bold=False,
          colour=INK2, align=PP_ALIGN.LEFT)
    for x in xs:
        V.node(slide, x, ny, nw, nh, "node")
    mid = ny + nh // 2
    for a, b in ((0, 1), (1, 2)):
        arrow(slide, xs[a] + nw, mid, xs[b], mid, colour=DARK, pt=2.25, head=False)
        label(slide, xs[a] + nw, mid - 400000, xs[b] - xs[a] - nw, 340000, "16379",
              size=14, bold=False, colour=INK)
    low = ny + nh + 450000
    arrow(slide, cx[0], ny + nh, cx[0], low, colour=DARK, pt=2.25, head=False)
    arrow(slide, cx[0], low, cx[2], low, colour=DARK, pt=2.25, head=False)
    arrow(slide, cx[2], low, cx[2], ny + nh, colour=DARK, pt=2.25, head=False)
    label(slide, cx[1] - 500000, low + 40000, 1000000, 340000, "16379", size=14,
          bold=False, colour=INK)
    notes(slide, "Clients talk to every node on the client port. Nodes talk to each other on the cluster bus, over a different port. The nodes exchange messages between them. It blindly trusts all incoming traffic on this port, there is no application level authentication over the link. It’s the operator’s responsibility to secure it. The connections are to be secured behind private networks or use mutual TLS.")
    return slide


def slide_embargo(prs):
    """Default embargo cut from 30 days to 14."""
    slide = titled(prs, "Embargoes are now two weeks")
    full_w = GRID_W - 1800000
    bh, gap = 800000, 450000
    top = MID_Y - (2 * bh + gap) // 2 + 150000
    for k, (who, days, f, fg) in enumerate((("before", 30, GREY, INK2),
                                            ("now", 14, AI, WHITE))):
        y = top + k * (bh + gap)
        label(slide, GRID_L, y, 1600000, bh, who, size=20, colour=INK2,
              align=PP_ALIGN.LEFT)
        V.chip(slide, GRID_L + 1800000, y, int(full_w * days / 30), bh, f"{days} days",
               fill=f, fg=fg, size=20)
    notes(slide, "The default embargo used to be 30 days, that’s too long in the current generation. As I mentioned 3 requests had the same vulnerability discovered all within two weeks. So, we need to get the patch out sooner to users to secure their system")
    return slide


def slide_security_group(prs):
    """Policy change: the security group now reaches beyond the TSC."""
    slide = titled(prs, "Security is no longer just the TSC")
    ow, oh = 8000000, 3400000
    ox, oy = (SLIDE_W - ow) // 2, MID_Y - oh // 2 + 200000
    rect(slide, ox, oy, ow, oh, fill=LILAC)
    label(slide, ox, oy + 120000, ow, 420000, "security group", size=20, colour=INK)
    iw, ih = 2600000, 1200000
    ix, iy = (SLIDE_W - iw) // 2, oy + (oh - ih) // 2 + 150000
    V.chip(slide, ix, iy, iw, ih, "TSC", fill=DARK, fg=WHITE, size=24)
    side = 1000000
    for cx in (ox + 300000, ox + 1450000, ox + ow - 300000 - side, ox + ow - 1450000 - side):
        icon(slide, "person", cx, iy + (ih - side) // 2, side)
    notes(slide, "And to get more hands to help with some of the steps after finding the bugs, we have expanded the trusted set of people beyond the maintainers and formed a new security group.")
    return slide


def slide_tweet(prs):
    slide = titled(prs, "Your next disclosure might be a tweet")
    path = os.path.join(ASSETS, "kimi-k3-redis-0day-tweet.png")
    from PIL import Image
    w_px, h_px = Image.open(path).size
    h = 4750000
    w = int(h * w_px / h_px)
    slide.shapes.add_picture(path, Emu((SLIDE_W - w) // 2), Emu(1080000), Emu(w), Emu(h))
    notes(slide, "Your next security vulnerability could be just a tweet, It might not go through the security workflow you had designed.\nA security researcher discovered a vulnerability using agents and published a proof of concept all under 30 minutes.\nBoth the attackers and defenders updated at the same time.")
    return slide


def slide_triage(prs):
    """Every incoming report goes through AI triage before a maintainer sees it."""
    slide = titled(prs, "AI triages every report first")
    rw, rh = 2300000, 900000
    ry = MID_Y - rh // 2 + 150000
    icon(slide, "report", GRID_L + (rw - 700000) // 2, ry - 800000, 700000)
    V.chip(slide, GRID_L, ry, rw, rh, "incoming report", fill=LILAC, fg=INK, size=16)
    tx, tw = GRID_L + rw + 700000, 2500000
    V.chip(slide, tx, ry, tw, rh, "AI triage", fill=AI, fg=WHITE, size=18)
    arrow(slide, GRID_L + rw, ry + rh // 2, tx, ry + rh // 2, colour=INK2)
    ow, oh, pitch = 3000000, 620000, 900000
    ox = GRID_L + GRID_W - ow
    outs = [("advisory", HUMAN, INK), ("ordinary bug", GREY, INK),
            ("not reproducible", GREY, INK2)]
    oy0 = ry + rh // 2 - (pitch * (len(outs) - 1) + oh) // 2
    for k, (text, f, fg) in enumerate(outs):
        y = oy0 + k * pitch
        arrow(slide, tx + tw, ry + rh // 2, ox, y + oh // 2, colour=INK2,
              elbow=None if y + oh // 2 == ry + rh // 2 else 3)
        V.chip(slide, ox, y, ow, oh, text, fill=f, fg=fg, size=16)
    notes(slide, "Triage is codified as prompts, per the Valkey blog. AI reproduces "
                 "the report against the target build and sorts it with the Act 1 rule "
                 "before a maintainer spends time on it. Only advisories go to the "
                 "maintainer queue as security work.\n\n"
                 "TODO: confirm the triage steps and outcomes with how Valkey actually "
                 "runs it.")
    return slide


LAB_PROGRAMS = [
    ("anthropic-logo.png", "Cyber Verification Program",
     ["Vetted defenders get its strongest models", "portal.anthropic.com/programs/cvp/apply"]),
    ("openai-logo.png", "Trusted Access for Cyber",
     ["Verified access to cyber-capable models", "openai.com/index/trusted-access-for-cyber"]),
]


def slide_programs(prs):
    """Two cards, one per lab: its logo, its security program, the link."""
    from PIL import Image
    slide = titled(prs, "Both labs run security programs for defenders")
    col_w, xs = box_columns(2)
    h = 2900000
    top = (SLIDE_H - h) // 2 + 200000
    max_w, max_h = col_w - 2 * PAD, 520000
    for x, (logo_file, program, lines) in zip(xs, LAB_PROGRAMS):
        rect(slide, x, top, col_w, h, fill=GREY)
        path = os.path.join(ASSETS, logo_file)
        w_px, h_px = Image.open(path).size
        logo_w = 3000000                                # so neither wordmark dominates
        lw, lh = logo_w, int(logo_w * h_px / w_px)     # fit the slot, keep the aspect
        if lh > max_h:
            lh, lw = max_h, int(max_h * w_px / h_px)
        slide.shapes.add_picture(path, Emu(x + PAD), Emu(top + PAD + (max_h - lh) // 2),
                                 Emu(lw), Emu(lh))
        hy = top + PAD + max_h + 300000
        text_box(slide, x + PAD, hy, max_w, 400000, [(program, 20, True, INK)])
        box = slide.shapes.add_textbox(Emu(x + PAD), Emu(hy + 500000), Emu(max_w),
                                       Emu(h - (hy + 500000 - top) - PAD))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for k, line in enumerate(lines):
            para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            para.space_after = Pt(6)
            r = para.add_run()
            r.text = line
            r.font.name, r.font.size = FONT, Pt(16)
            r.font.color.theme_color = AI if k == len(lines) - 1 else INK
    notes(slide, "Anthropic: the Cyber Verification Program, expanded 6 October 2026. For "
                 "security teams, critical infrastructure operators, open-source maintainers, "
                 "red teams and pen-testers. Vetted defenders get Opus 5.5, Sonnet 5.5 and "
                 "Mythos 5.1 with fewer cyber blocks, in three tiers: Defense Access, Red "
                 "Team Access, Specialized Access. anthropic.com/news/cyber-verification-"
                 "program\n\n"
                 "OpenAI: Trusted Access for Cyber, 5 February 2026. An identity and "
                 "trust-based framework for access to its most cyber-capable models, with "
                 "$10M in API credits and the Cybersecurity Grant Program.")
    return slide


BACKPORTS = [("9.1", "17:03"), ("7.2", "17:04"), ("9.0", "17:11"), ("8.1", "17:12"),
             ("8.0", "17:15")]


def slide_backport_timeline(prs):
    """The TLS fix fanning out from unstable to five release branches, with merge times."""
    slide = titled(prs, "Five release branches in 72 minutes")
    sw, sh = 2700000, 1000000
    bw, bh, pitch = 2600000, 560000, 760000
    top = 1500000
    sy = top + (pitch * (len(BACKPORTS) - 1) + bh) // 2 - sh // 2
    V.chip(slide, GRID_L, sy, sw, sh, "unstable  ·  16:03", fill=DARK, fg=WHITE, size=18)
    bot_x = GRID_L + sw + 900000
    # Above the arrow, never on it.
    icon(slide, "model", bot_x, sy + sh // 2 - 1000000, 560000)
    label(slide, bot_x - 400000, sy + sh // 2 - 420000, 1360000, 340000, "backport sweep",
          size=14, bold=False, colour=INK2)
    bx = GRID_L + GRID_W - bw
    for i, (branch, when) in enumerate(BACKPORTS):
        y = top + i * pitch
        arrow(slide, GRID_L + sw, sy + sh // 2, bx, y + bh // 2, colour=INK2,
              elbow=None if y + bh // 2 == sy + sh // 2 else 3)
        V.chip(slide, bx, y, bw, bh, f"{branch}  ·  {when}", fill=AI, fg=WHITE, size=16)
    source_note(slide, "github.com/valkey-io/valkey, PRs #4234 to #4242, 21 July 2026, UTC")
    notes(slide, "The TLS fix, #4234, merged to unstable at 16:03 UTC on 21 July. The "
                 "valkey-ci-agent bot opened backport sweeps from 16:05. All five release "
                 "branches had it merged by 17:15. The same bot opened release-note PRs "
                 "for 7.2.14, 8.0.10, 8.1.9, 9.0.5 and 9.1.1, all merged by 23:39. One fix, "
                 "five releases, one day.\n\n"
                 "7.2 needed an adapted fix: it lacks a flag the fix uses. The bot kept the "
                 "security change, dropped the part that didn't apply, and flagged it for "
                 "review on #4242.")
    return slide


# (step, UTC time, who does it) for the TLS fix on 21 July 2026. "bot" steps are the
# valkey-ci-agent backport and release-notes sweeps and the valkey-release-automation
# post-release jobs; "human" steps are maintainers merging and publishing.
RELEASE_FLOW = [
    ("fix merged", "16:03", "human"),
    ("backport sweep", "16:05", "bot"),
    ("backports merged", "17:15", "human"),
    ("release notes", "17:23", "bot"),
    ("releases published", "23:41", "human"),
    ("downstream updates", "00:08", "bot"),          # hashes, containers, website
]


def slide_release_flow(prs):
    """From fix to five published releases: bot steps blue, human steps rose."""
    slide = titled(prs, "Fix to five published releases in under eight hours")
    gap, h = 100000, 1900000
    w = (GRID_W - gap * (len(RELEASE_FLOW) - 1)) // len(RELEASE_FLOW)
    y = MID_Y - h // 2 + 150000
    for i, (step, when, who) in enumerate(RELEASE_FLOW):
        x = GRID_L + i * (w + gap)
        bg, fg = (AI, WHITE) if who == "bot" else (HUMAN, INK)
        rect(slide, x, y, w, h, fill=bg)
        if who == "human":                  # no icon on blue: it vanishes there
            icon(slide, "person", x + (w - 520000) // 2, y + 200000, 520000, on=bg)
        label(slide, x + 60000, y + 820000, w - 120000, 600000, step, size=15, colour=fg)
        label(slide, x, y + 1400000, w, 380000, when, size=18, colour=fg)
    source_note(slide, "valkey-io/valkey, valkey-hashes, valkey-container, "
                       "valkey-io.github.io, 21 to 22 July 2026, UTC")
    notes(slide, "Blue is automation, rose is a person. Bots opened the backport sweeps and "
                 "the release-note PRs for 7.2.14, 8.0.10, 8.1.9, 9.0.5 and 9.1.1. A "
                 "maintainer published the five releases between 21:35 and 23:41. Each "
                 "publish triggered the release automation: hashes updated within a minute, "
                 "and container and website PRs opened, all merged by 00:08 on 22 July.\n\n"
                 "Fix merged to five releases published: 7 hours 38 minutes.")
    return slide


def slide_getting_involved(prs):
    slide = on_layout(prs, V.LAYOUT_LIST_ROWS_IMAGE)
    set_title(slide, "Getting involved")
    list_rows(slide, ["Read the post: valkey.io/blog",
                      "Report: security@lists.valkey.io",
                      "Harden: valkey.io/blog/properly-secure-your-valkey-deployment"])
    slide.placeholders[V.LIST_ROWS + 1].insert_picture(
        os.path.join(ASSETS, "cve-tsunami-blog-qr.png"))
    notes(slide, "The QR goes to the security post, which links the hardening guide.")
    return slide


# --- the deck ------------------------------------------------------------------


def build(out=OUT):
    tmp = out + ".building"
    shutil.copyfile(TEMPLATE, tmp)
    prs = Presentation(tmp)
    V.strip_slides(prs)

    starts = {}
    # Opening: the problem
    starts["Opening"] = len(prs.slides)
    # Two presenters and the template's default band art (the security shield).
    notes(V.title_band(prs, ["Open-Source Maintainers", "vs. The CVE Tsunami"],
                       [("Madelyn Olson", "Valkey maintainer"),
                        ("Harkrishn Patro", "Valkey maintainer")]),
          "This year has been very interesting for all the open source projects in terms of security. Tsunami of CVE requests hitting the maintainers constantly. I, Harkrishn Patro, Valkey maintainer joined by Madelyn Olson, a Valkey maintainer as well. We’re here to talk about how we went about the onslaught of CVEs and came out of it stronger, better, faster.")

    slide_tweet(prs)                    # the cold open
    slide_advisories(prs)
    # 32pt, not the layout's 40pt: the verbatim quote is 120 characters and overflows
    # the statement box at 40.
    notes(statement(prs, "“a seemingly unending deluge of vulnerability reports "
                         "from researchers using LLMs to mine for security gold in our "
                         "software”",
                    "Jeremy Stanley, OpenStack, oss-security, April 2026", size=32),
          "This has been the theme of this year. A never ending stream of CVE requests by security researchers.")
    slide_three_reporters(prs)          # after the quote, per Hari's order
    slide_pipeline(prs)

    # Act 1
    starts["Act 1: CVEs aren't so special anymore"] = len(prs.slides)
    statement(prs, "Act 1: CVEs aren't so special anymore")
    notes(statement(prs, "“When everything is a CVE, nothing is”",
                    "-Madelyn Olson (Wise Maintainer)"),
          "The thesis of the act. If every crash gets a CVE, the ones that matter drown in "
          "the noise, for us and for every operator reading advisories.")
    slide_acl(prs)
    notes(cards(prs, "Raising the bar on CVEs",
                [("security", "Advisory",
                  ["Pre-auth crash", "Authenticated out-of-bounds read",
                   "Memory corruption with a path to code execution"]),
                 ("reliability", "Ordinary bug",
                  ["Authenticated crash or hang", "Crash from a client with EVAL"])],
                (HUMAN, LILAC), h=3300000),
          "Prior to this surge anything touching availability, confidentiality, or integrity got an advisory. We are more careful going forward with assigning a CVE to a vulnerability.")
    slide_cluster_bus(prs)
    notes(cards(prs, "The cluster bus gets hardening instead of CVEs",
                [("security", "No CVE", ["Reaching the bus already means owning the cluster"]),
                 ("hardening", "Still hardened",
                  ["Bus parsing bugs still get fixed", "Mutual TLS with tls-cluster yes",
                   "Firewall guidance in the docs"])],
                (HUMAN, LILAC), h=3000000),
          "An attacker who can reach the bus can already hijack slots, so a crash there gives them nothing new.")
    slide_embargo(prs)
    slide_security_group(prs)

    # Act 2
    starts["Act 2: Find bugs before they do"] = len(prs.slides)
    statement(prs, "Act 2: Find bugs before they do")
    for f in (1, 2, 3, 4):
        slide_audit_funnel(prs, f)
    notes(big_number(prs, "34", "real bugs: Valkey, JSON, search, bloom", BLOG),
          "Audits run periodically on several frontier models. Goal is before every "
          "release.")
    slide_tls_code(prs)                 # the code first, then the build explains it
    for f in (1, 2, 3, 4):
        slide_tls_list(prs, f)
    slide_what_about_rust(prs)
    slide_rust(prs, 1)
    slide_rust(prs, 2)
    slide_rust_rows(prs)

    # Act 3
    starts["Act 3: React fast"] = len(prs.slides)
    statement(prs, "Act 3: React fast")
    slide_triage(prs)
    slide_programs(prs)
    notes(cards(prs, "Same-day fixes on every branch",
                [("person", "A year ago", ["Hand cherry-pick into every branch"]),
                 ("model", "Now", ["AI backport PRs, conflicts resolved"])],
                (HUMAN, LILAC), h=2700000),
          "Triage and release steps are codified as prompts. Goal is a zero-day fix "
          "shipped the same day.")
    slide_backport_timeline(prs)
    slide_release_flow(prs)
    notes(statement(prs, "Humans review the changes."),
          "Automation writes the backport. A person is accountable for it, more so in "
          "a security release.")

    # Close
    starts["Close"] = len(prs.slides)
    # Recap: one card per act, the act and its takeaway.
    notes(cards(prs, "Recap",
                [("security", "Act 1", ["CVEs aren't so special anymore",
                                        "Advisories are for crossing a boundary"]),
                 ("hardening", "Act 2", ["Find bugs before they do",
                                         "Adversarial audits that end in a crashing test"]),
                 ("migration", "Act 3", ["React fast",
                                         "Same-day backports that humans review"])],
                (GREY,) * 3, h=3000000, bold_first=True),
          "Operators: stay on current patches and make consuming updates routine; follow "
          "the hardening guide. Researchers: email security@lists.valkey.io with a "
          "reproducer against the target build, ideally with a fix. Maintainers: build "
          "the harness, write the policy, automate the backports.")
    notes(statement(prs, "Assume bug-finding keeps getting cheaper.", bold=True),
          "Plan for it. Every process in this talk is built on that assumption. Getting "
          "these tools to maintainers is what keeps open source secure.")
    slide_getting_involved(prs)
    s = V.thank_you(prs, "Thank you! Let’s connect.",
                    [("Madelyn Olson", ["@reconditerose (X and Bluesky)", "madelynolson.com",
                                        "linkedin.com/in/madelyn-olson-valkey"]),
                     ("Harkrishn Patro", ["@buntyhp (X)",
                                          "linkedin.com/in/harkrishn-patro"])])

    # PowerPoint sections match the acts. Last, once the order is final.
    V.set_sections(prs, list(starts.items()))

    # Font pass last, over the finished deck. Code runs keep Courier New.
    for old in ("Play", "Century Gothic"):
        retypeface(prs, old, FONT)
    for slide in prs.slides:
        for sh in shapes_flat(slide):
            for para in sh.text_frame.paragraphs if sh.has_text_frame else []:
                for run in para.runs:
                    if run.font.name != CODE_FONT:
                        run.font.name = FONT
    audit(prs)

    prs.save(tmp)
    Presentation(tmp)                   # reopen: if it raises, PowerPoint would too
    os.replace(tmp, out)
    print(f"wrote {out}: {len(prs.slides)} slides")
    return out


def audit(prs):
    audit_paragraph_order(prs)
    V.audit_sections(prs)
    # A title makes one point. A comma hedges it or adds a second one.
    for i, slide in enumerate(prs.slides, 1):
        for ph in slide.placeholders:
            idx = ph.placeholder_format.idx
            is_title = idx == 0 and slide.slide_layout.name != V.LAYOUT_TITLE_BAND
            is_statement = idx == 1 and slide.slide_layout.name in (
                V.LAYOUT_INSIGHT, V.LAYOUT_INSIGHT_BOLD)
            if not (is_title or is_statement):
                continue
            text = ph.text_frame.text
            if not text.startswith("“") and "," in text:       # quotes stay verbatim
                raise SystemExit(f"slide {i}: comma in title {ph.text_frame.text!r}")
    from pptx.oxml.ns import qn
    for i, slide in enumerate(prs.slides, 1):
        for ph in slide.placeholders:
            xfrm = ph._element.spPr.find(qn("a:xfrm"))
            if xfrm is not None and (xfrm.find(qn("a:off")) is None
                                     or xfrm.find(qn("a:ext")) is None):
                raise SystemExit(f"slide {i}: placeholder {ph.name!r} has a partial "
                                 "xfrm; PowerPoint will misplace it")
    audit_list_rows(prs, V.LAYOUT_LIST_ROWS, V.LAYOUT_LIST_ROWS_IMAGE)
    for i, slide in enumerate(prs.slides, 1):
        ids = [sh.shape_id for sh in shapes_flat(slide)]
        if len(ids) != len(set(ids)):
            raise SystemExit(f"slide {i}: duplicate shape ids")
        for sh in shapes_flat(slide):
            if sh.left is None:
                continue
            if sh.left < 0 or sh.top < 0 or sh.left + sh.width > SLIDE_W \
                    or sh.top + sh.height > SLIDE_H:
                raise SystemExit(f"slide {i}: {sh.name!r} out of bounds")
            for para in sh.text_frame.paragraphs if sh.has_text_frame else []:
                for run in para.runs:
                    if run.font.name not in (FONT, CODE_FONT):
                        raise SystemExit(f"slide {i}: run in {run.font.name!r}")
    for part in prs.part.package.iter_parts():
        if str(part.partname).endswith(".xml"):
            blob = part.blob.decode("utf-8", "replace")
            for bad in ("Play", "Century Gothic"):
                if f'typeface="{bad}"' in blob:
                    raise SystemExit(f"{bad!r} still in {part.partname}")


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else OUT)
